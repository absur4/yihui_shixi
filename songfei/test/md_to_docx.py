"""把《汇报稿》Markdown 转成 Word，只保留每页 PPT 的演讲内容。

规则：
- 只取标题形如 `## 第 N 页 · …` 的章节，其余（时间分配总表、可删段落、附录）全部丢弃
- 每页正文取行首为 `>` 的行（即"口播"），一行一段
- 标题与"**口播**"之间、那些 `**看图上什么**：…` 之类的提示行，压缩成一行灰色小字
  放在标题下，方便临场看但不干扰朗读；不需要时用 --no-cues 去掉
- 每页 PPT 单独起一页，方便跟着翻页

用法：
    python md_to_docx.py                     # 默认转 汇报稿_15分钟.md
    python md_to_docx.py --no-cues           # 连提示行也不要，只留演讲词
    python md_to_docx.py --input a.md --output b.docx
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PAGE_HEADING = re.compile(r"^第\s*\d+\s*页")
BOLD = re.compile(r"\*\*(.+?)\*\*")
CODE = re.compile(r"`([^`]+)`")

FONT = "微软雅黑"
HEADING_COLOR = RGBColor(0x1F, 0x3B, 0x63)
CUE_COLOR = RGBColor(0x8A, 0x8A, 0x8A)


def set_font(run, size=None, color=None, bold=False, italic=False):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attribute in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(attribute), FONT)


def add_text(paragraph, text, size, color=None, bold=False, italic=False):
    """写入一段文字，并把 **粗体** 与 `代码` 的内联标记处理掉。"""
    position = 0
    for match in BOLD.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(CODE.sub(r"\1", text[position:match.start()]))
            set_font(run, size, color, bold, italic)
        run = paragraph.add_run(CODE.sub(r"\1", match.group(1)))
        set_font(run, size, color, bold=True, italic=italic)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(CODE.sub(r"\1", text[position:]))
        set_font(run, size, color, bold, italic)


def parse(markdown: str):
    """→ [(标题, 提示行列表, 演讲段落列表)]，按出现顺序。"""
    slides = []
    current = None
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            title = line[3:].strip()
            current = (title, [], []) if PAGE_HEADING.match(title) else None
            if current:
                slides.append(current)
            continue
        if current is None:
            continue
        title, cues, spoken = current
        if line.startswith(">"):
            text = line[1:].strip()
            if text:
                spoken.append(text)
        elif line and not line.startswith(("---", "#", "|")):
            if line.strip() != "**口播**":
                cues.append(line.strip())
    return slides


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, end):
        run._r.append(node)
    set_font(run, 9, CUE_COLOR)


def main() -> int:
    parser = argparse.ArgumentParser(description="汇报稿 Markdown → Word")
    parser.add_argument("--input", default=str(HERE / "汇报稿_15分钟.md"))
    parser.add_argument("--output", default=str(HERE / "汇报稿_15分钟.docx"))
    parser.add_argument("--no-cues", action="store_true",
                        help="不要提示行，只留演讲词")
    args = parser.parse_args()

    source = Path(args.input)
    slides = parse(source.read_text(encoding="utf-8"))
    if not slides:
        raise SystemExit(f"没解析出任何『第 N 页』章节：{source}")

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(14)
    add_text(heading, f"四种通信中间件性能对比测试 · 汇报稿（15 分钟，共 {len(slides)} 页）",
             15, HEADING_COLOR, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(footer)

    for index, (title, cues, spoken) in enumerate(slides):
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        block = document.add_paragraph()
        block.paragraph_format.space_after = Pt(6)
        add_text(block, title, 14, HEADING_COLOR, bold=True)

        if cues and not args.no_cues:
            cue = document.add_paragraph()
            cue.paragraph_format.space_after = Pt(12)
            add_text(cue, "  ".join(cues), 9, CUE_COLOR, italic=True)

        for paragraph_text in spoken:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(9)
            paragraph.paragraph_format.line_spacing = 1.5
            add_text(paragraph, paragraph_text, 12)

    target = Path(args.output)
    document.save(str(target))
    print(f"[Word] {target}")
    print(f"  页数（PPT 页）：{len(slides)}   演讲段落：{sum(len(s[2]) for s in slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
